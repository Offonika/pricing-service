from __future__ import annotations

import json
import re
import shutil
import urllib.request
import zipfile
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import HTTPException
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from sqlalchemy import text
from sqlalchemy.engine import Engine

from app.core.config import Settings, get_settings
from app.infrastructure.db import build_onec_engine_from_settings
from app.schemas.procurement_labels import (
    ProcurementCertificationDocsGenerateResponse,
    ProcurementLabelGenerateResponse,
    ProcurementLabelOrderPreview,
    ProcurementLabelRow,
)
from app.services.bank_payments_bitrix import BitrixDiskClient, BitrixDiskError

LABEL_WIDTH_MM = 58
LABEL_HEIGHT_MM = 40
LABEL_DPI = 300
LABEL_WIDTH_PX = round(LABEL_WIDTH_MM / 25.4 * LABEL_DPI)
LABEL_HEIGHT_PX = round(LABEL_HEIGHT_MM / 25.4 * LABEL_DPI)
BLOCK_NON_VED = "Карточка не в контуре ВЭД импорт"
BLOCK_MISSING_ONEC = "Не найден 1С-код товара"
BLOCK_MISSING_SKU = "Не найден SKU/артикул товара"
BLOCK_MISSING_BARCODE = "Не найден штрихкод 1С"
BLOCK_MISSING_CERTIFICATE = "Нет связки товара с сертификатом/ДС"
BLOCK_EXPIRED_CERTIFICATE = "ДС/сертификат истек"
BLOCK_MISSING_CERTIFICATE_FILE = "У ДС/сертификата не прикреплен файл"
BLOCK_AMBIGUOUS_CERTIFICATE = "Найдено несколько активных ДС/сертификатов для SKU"
BLOCK_UNVERIFIED_CERTIFICATE = "ДС/сертификат не подтвержден"
BLOCK_MISSING_CERTIFICATION_ROWS = "Не найдены строки заказа для пакета сертификации"

CERTIFICATE_ENTITY_TYPE_ID = 1052
PRODUCT_PASSPORT_ENTITY_TYPE_ID = 1044
ACTIVE_CERTIFICATE_STATUSES = {"covered", "verified"}
WARNING_MISSING_PRODUCT_PASSPORT = "Не найден паспорт товара ВЭД"
WARNING_MISSING_TRADE_NAME = "В паспорте товара нет Trade name"
WARNING_MISSING_TNVED = "В паспорте товара нет ТН ВЭД"

CERTIFICATION_REQUIRED_FIELDS = (
    "sku_1c",
    "trade_name_family",
    "compatibility",
    "capacity_mah",
    "voltage_v",
    "energy_wh",
    "dim",
    "bms",
    "connector",
    "tnved",
    "un_code",
    "gtin_ean13",
    "gost_r_ds_number",
    "gost_r_ds_covers_sku",
)

CERTIFICATION_DOCUMENT_CHECKLIST = [
    "Заявитель: карточка компании, ИНН/ОГРН, реквизиты, ЭЦП для ФГИС/реестра.",
    "Партия: заказ поставщику, инвойс, packing list, контракт или спецификация.",
    "Товар: мастер-реестр SKU, TradeName/family_formula, совместимость, mAh, V, Wh, размеры, BMS, connector.",
    "Классификация: ТН ВЭД, UN Code, страна происхождения, производитель и поставщик.",
    "АКБ-документы: UN38.3, MSDS/SDS, battery statement, RoHS/CE-EMC при наличии.",
    "Доказательная база: протоколы испытаний или запрос на испытания по выбранной схеме.",
    "Маркировка: фото/макет этикетки, упаковка, EAC/РСТ только после подтверждения права.",
    "GTIN/EAN-13: заказать на SKU, где код отсутствует, и вернуть номера в мастер-реестр.",
]

CERTIFICATION_SHEET_TITLES = {
    "master": "Мастер-реестр",
    "gtin": "Заказ GTIN",
}

CERTIFICATION_COLUMN_TITLES = {
    "line_no": "№ строки",
    "onec_order": "Заказ 1С",
    "onec_item_code": "Код номенклатуры 1С",
    "article_1c": "Артикул 1С (числовой)",
    "sku_1c": "SKU 1С (отдельное поле)",
    "item_name_1c": "Наименование 1С",
    "qty": "Количество",
    "unit": "Ед.",
    "barcode_1c": "Штрихкод 1С",
    "barcode_source": "Источник штрихкода",
    "generated_sku": "SKU генератора (кандидат)",
    "generation_status": "Статус генерации SKU",
    "external_sku_candidate": "Старый F5/OEM SKU (кандидат)",
    "trade_name_family": "TradeName / семейная формула для ДС",
    "compatibility": "Совместимость",
    "capacity_mah": "Ёмкость, mAh",
    "voltage_v": "Напряжение, V",
    "energy_wh": "Энергия, Wh",
    "dim": "Размеры, DIM",
    "bms": "BMS / плата защиты",
    "connector": "Коннектор",
    "tnved": "ТН ВЭД",
    "un_code": "UN Code",
    "current_eaeu_ds_number": "Текущая ДС ЕАЭС",
    "gost_r_ds_number": "Новая ДС ГОСТ Р",
    "new_gost_r_ds_number": "Новая ДС ГОСТ Р",
    "gost_r_ds_covers_sku": "Покрытие новой ДС по SKU",
    "gtin_ean13": "GTIN/EAN-13",
    "gtin_action": "Действие по GTIN",
    "missing_fields": "Что не заполнено / ошибки",
    "review_status": "Статус проверки",
    "next_action": "Что сделать",
}

CERTIFICATION_REVIEW_STATUS_TITLES = {
    "missing_sku_1c": "Нет SKU 1С",
    "needs_gtin_and_gost_r_ds": "Нет GTIN и новой ДС ГОСТ Р",
    "needs_gost_r_ds": "Нет новой ДС ГОСТ Р / нет покрытия SKU",
    "needs_gtin": "Нужен GTIN/EAN-13",
    "needs_product_properties": "Не заполнены свойства товара",
    "ready_for_certifier": "Готово к проверке сертификатором",
}

CERTIFICATION_BARCODE_SOURCE_TITLES = {
    "1c_internal": "Штрихкод из 1С",
    "catalog_barcode": "Штрихкод из каталога",
    "catalog_gtin": "GTIN из каталога",
}

CERTIFICATION_MASTER_FILENAME = "мастер-таблица-для-декларации.xlsx"
LEGACY_CERTIFICATION_MASTER_FILENAME = "master-register.xlsx"

EAN_LEFT_ODD = {
    "0": "0001101",
    "1": "0011001",
    "2": "0010011",
    "3": "0111101",
    "4": "0100011",
    "5": "0110001",
    "6": "0101111",
    "7": "0111011",
    "8": "0110111",
    "9": "0001011",
}
EAN_LEFT_EVEN = {
    "0": "0100111",
    "1": "0110011",
    "2": "0011011",
    "3": "0100001",
    "4": "0011101",
    "5": "0111001",
    "6": "0000101",
    "7": "0010001",
    "8": "0001001",
    "9": "0010111",
}
EAN_RIGHT = {
    "0": "1110010",
    "1": "1100110",
    "2": "1101100",
    "3": "1000010",
    "4": "1011100",
    "5": "1001110",
    "6": "1010000",
    "7": "1000100",
    "8": "1001000",
    "9": "1110100",
}
EAN_PARITY = {
    "0": "OOOOOO",
    "1": "OOEOEE",
    "2": "OOEEOE",
    "3": "OOEEEO",
    "4": "OEOOEE",
    "5": "OEEOOE",
    "6": "OEEEOO",
    "7": "OEOEOE",
    "8": "OEOEEO",
    "9": "OEEOEO",
}


@dataclass(frozen=True)
class BitrixFileResult:
    file_id: str
    url: str


@dataclass(frozen=True)
class CertificateCoverage:
    item_id: str = ""
    certificate_id: str = ""
    certificate_number: str = ""
    status: str = "missing"
    valid_to: str = ""
    file_id: str = ""
    eac_allowed: bool = False
    covered_skus: tuple[str, ...] = ()
    source: str = "catalog"


@dataclass(frozen=True)
class ProductPassport:
    item_id: str = ""
    sku: str = ""
    onec_item_code: str = ""
    trade_name: str = ""
    tnved: str = ""
    manufacturer: str = ""
    product_series: str = ""


@dataclass(frozen=True)
class CertificateResolution:
    status: str = "missing"
    certificate: CertificateCoverage | None = None
    blockers: tuple[str, ...] = ()


@dataclass(frozen=True)
class ExistingLabelState:
    status: str = ""
    version: int | None = None
    zip_url: str = ""
    disk_file_id: str = ""


@dataclass(frozen=True)
class ExistingCertificationDocsState:
    status: str = ""
    version: int | None = None
    zip_url: str = ""
    disk_file_id: str = ""


@dataclass(frozen=True)
class BitrixLabelSources:
    certificate_catalog: dict[str, Any] = field(default_factory=dict)
    product_passports: dict[str, ProductPassport] = field(default_factory=dict)


class ProcurementLabelsBitrixClient:
    def __init__(
        self,
        webhook_url: str,
        *,
        timeout: float = 10.0,
        urlopen: Any = urllib.request.urlopen,
    ) -> None:
        self.webhook_url = webhook_url.rstrip("/")
        self.timeout = timeout
        self._urlopen = urlopen

    def call(self, method: str, params: dict[str, Any] | None = None) -> dict[str, Any]:
        request = urllib.request.Request(
            f"{self.webhook_url}/{method}.json",
            data=json.dumps(params or {}, ensure_ascii=False).encode("utf-8"),
            headers={"Content-Type": "application/json", "Accept": "application/json"},
            method="POST",
        )
        with self._urlopen(request, timeout=self.timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
        if payload.get("error"):
            message = (
                f"Bitrix24 {method}: {payload['error']} " f"{payload.get('error_description', '')}"
            ).strip()
            raise RuntimeError(message)
        return payload

    def get_item(self, *, entity_type_id: int, item_id: str) -> dict[str, Any]:
        payload = self.call("crm.item.get", {"entityTypeId": entity_type_id, "id": item_id})
        result = payload.get("result") or {}
        item = result.get("item") if isinstance(result, dict) else {}
        if not isinstance(item, dict):
            raise RuntimeError("Bitrix24 crm.item.get returned invalid payload")
        return item

    def list_items(
        self,
        *,
        entity_type_id: int,
        select: list[str] | None = None,
        filter_: dict[str, Any] | None = None,
        limit: int = 1000,
    ) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        start: int | None = 0
        while start is not None and len(rows) < limit:
            params: dict[str, Any] = {
                "entityTypeId": entity_type_id,
                "order": {"id": "ASC"},
                "start": start,
            }
            if select:
                params["select"] = select
            if filter_:
                params["filter"] = filter_
            payload = self.call("crm.item.list", params)
            result = payload.get("result") or {}
            items = result.get("items") if isinstance(result, dict) else []
            if not isinstance(items, list):
                raise RuntimeError("Bitrix24 crm.item.list returned invalid payload")
            rows.extend(item for item in items if isinstance(item, dict))
            next_start = payload.get("next")
            start = int(next_start) if next_start is not None else None
        return rows[:limit]

    def update_item(self, *, entity_type_id: int, item_id: str, fields: dict[str, Any]) -> None:
        if not fields:
            return
        self.call(
            "crm.item.update",
            {"entityTypeId": entity_type_id, "id": item_id, "fields": fields},
        )


def clean_string(value: Any) -> str:
    return str(value or "").strip()


def normalize_lookup_key(value: Any) -> str:
    return clean_string(value).upper()


def split_sku_list(value: Any) -> tuple[str, ...]:
    raw = clean_string(value)
    if not raw:
        return ()
    parts = [part.strip() for part in re.split(r"[\n,;\t ]+", raw) if part.strip()]
    seen: set[str] = set()
    result: list[str] = []
    for part in parts:
        key = normalize_lookup_key(part)
        if key and key not in seen:
            seen.add(key)
            result.append(part)
    return tuple(result)


def _bool_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    token = clean_string(value).casefold()
    return token in {"1", "y", "yes", "true", "да", "истина"}


def _date_from_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = clean_string(value)
    if not raw:
        return None
    normalized = raw.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(normalized).date()
    except ValueError:
        pass
    try:
        return date.fromisoformat(raw[:10])
    except ValueError:
        return None


def _is_expired(value: Any, *, today: date | None = None) -> bool:
    valid_to = _date_from_value(value)
    if valid_to is None:
        return False
    return valid_to < (today or date.today())


def _file_id(value: Any) -> str:
    if isinstance(value, dict):
        return clean_string(value.get("id") or value.get("ID") or value.get("fileId"))
    if isinstance(value, list):
        for item in value:
            resolved = _file_id(item)
            if resolved:
                return resolved
    return clean_string(value)


def load_json_file(path: Path) -> Any:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def load_mapping(settings: Settings | None = None) -> dict[str, Any]:
    settings = settings or get_settings()
    path = Path(settings.procurement_labels_mapping_path)
    payload = load_json_file(path)
    return payload if isinstance(payload, dict) else {}


def crm_item_rest_field_name(field: str) -> str:
    raw = clean_string(field)
    if not raw:
        return ""
    builtins = {
        "ID": "id",
        "TITLE": "title",
        "STAGE_ID": "stageId",
        "CATEGORY_ID": "categoryId",
        "ASSIGNED_BY_ID": "assignedById",
    }
    upper = raw.upper()
    if upper in builtins:
        return builtins[upper]
    if upper.startswith("UF_CRM_"):
        parts = [part for part in raw.split("_")[2:] if part]
        return "ufCrm" + "".join(part[:1].upper() + part[1:].lower() for part in parts)
    return raw


def item_value(item: dict[str, Any], mapping: dict[str, Any], logical_key: str) -> Any:
    field = ((mapping.get("field_map") or {}).get(logical_key)) or logical_key
    candidates = [
        str(field),
        crm_item_rest_field_name(str(field)),
        logical_key,
    ]
    for key in candidates:
        if key in item:
            return item[key]
    return None


def _decimal(value: Any) -> Decimal:
    if isinstance(value, Decimal):
        return value
    raw = clean_string(value).replace(" ", "").replace(",", ".")
    if not raw:
        return Decimal("0")
    return Decimal(raw)


def _int_or_none(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def source_number_candidates(number: str) -> list[str]:
    original = clean_string(number)
    if not original:
        return []
    candidates = [original]
    match = re.match(r"^(.*?)(\d+)$", original)
    if not match:
        return candidates
    prefix, digits = match.groups()
    numeric = int(digits)
    for width in (len(digits) - 1, len(digits) + 1):
        if width <= 0:
            continue
        candidate = prefix + str(numeric).zfill(width)
        if candidate not in candidates:
            candidates.append(candidate)
    return candidates


def normalize_contour(value: Any) -> str:
    token = clean_string(value).casefold().replace(" ", "").replace("_", "").replace("-", "")
    if token in {"vedimport", "вэдимпорт"}:
        return "ved_import"
    if token in {"cargo", "карго"}:
        return "cargo"
    if token in {"ordinary", "обычный", "обычная"}:
        return "ordinary"
    return token


def item_procurement_contour(item: dict[str, Any], mapping: dict[str, Any]) -> str:
    category_id = _int_or_none(item.get("categoryId") or item.get("CATEGORY_ID"))
    ved_category = _int_or_none(
        ((mapping.get("category_map") or {}).get("ved_import") or {}).get("id")
    )
    if category_id is not None and ved_category is not None and category_id == ved_category:
        return "ved_import"

    raw_value = item_value(item, mapping, "procurement_contour")
    enum_id = clean_string(
        ((mapping.get("enum_map") or {}).get("procurement_contour") or {}).get("ved_import")
    )
    if enum_id and clean_string(raw_value) == enum_id:
        return "ved_import"
    return normalize_contour(raw_value)


def item_onec_number(item: dict[str, Any], mapping: dict[str, Any]) -> str:
    value = clean_string(item_value(item, mapping, "onec_source_number"))
    if value:
        return value
    title = clean_string(item.get("title") or item.get("TITLE"))
    match = re.search(r"[А-ЯA-Z]{2,}\d{6,}", title)
    return match.group(0) if match else ""


def _nested_mapping(mapping: dict[str, Any], key: str, fallback: str = "") -> dict[str, Any]:
    value = mapping.get(key)
    if isinstance(value, dict):
        return value
    if fallback:
        value = mapping.get(fallback)
        if isinstance(value, dict):
            return value
    return {}


def _logical_field_value(
    item: dict[str, Any],
    field_map: dict[str, Any],
    logical_key: str,
    *fallback_fields: str,
) -> Any:
    field = clean_string(field_map.get(logical_key))
    candidates = [field, crm_item_rest_field_name(field), logical_key, *fallback_fields]
    for key in candidates:
        if key and key in item:
            return item[key]
    return None


def _enum_status(
    value: Any,
    enum_map: dict[str, Any],
    *,
    labels: dict[str, str] | None = None,
    default: str = "",
) -> str:
    raw = clean_string(value)
    if not raw:
        return default
    raw_token = raw.casefold()
    for logical_key, enum_id in enum_map.items():
        if raw == clean_string(enum_id) or raw_token == clean_string(logical_key).casefold():
            return str(logical_key)
    labels = labels or {}
    for label, logical_key in labels.items():
        if raw_token == label.casefold():
            return logical_key
    return raw


def _existing_label_state(item: dict[str, Any], mapping: dict[str, Any]) -> ExistingLabelState:
    field_map = mapping.get("field_map") or {}
    enum_map = (mapping.get("enum_map") or {}).get("label_generation_status") or {}
    status = _enum_status(
        item_value(item, mapping, "label_generation_status"),
        enum_map,
        labels={
            "Черновик": "draft",
            "Заблокировано": "blocked",
            "Утверждено": "approved",
            "Отправлено фабрике": "sent_to_factory",
        },
    )
    version = _int_or_none(item_value(item, mapping, "label_generation_version"))
    zip_url = clean_string(item_value(item, mapping, "label_generation_zip_url"))
    disk_file_id = clean_string(
        _logical_field_value(
            item,
            field_map,
            "label_generation_disk_file_id",
            "ufCrm8Labelgenerationdiskfileid",
        )
    )
    return ExistingLabelState(
        status=status,
        version=version,
        zip_url=zip_url,
        disk_file_id=disk_file_id,
    )


def _existing_certification_docs_state(
    item: dict[str, Any], mapping: dict[str, Any]
) -> ExistingCertificationDocsState:
    field_map = mapping.get("field_map") or {}
    enum_map = (mapping.get("enum_map") or {}).get("certification_docs_status") or {}
    status = _enum_status(
        item_value(item, mapping, "certification_docs_status"),
        enum_map,
        labels={
            "Черновик": "draft",
            "Нужны данные": "needs_data",
            "Заблокировано": "blocked",
            "Пакет готов": "ready",
            "GTIN заказан": "gtin_requested",
            "Передано сертификатору": "sent_to_certifier",
        },
    )
    version = _int_or_none(item_value(item, mapping, "certification_docs_version"))
    zip_url = clean_string(item_value(item, mapping, "certification_docs_zip_url"))
    disk_file_id = clean_string(
        _logical_field_value(
            item,
            field_map,
            "certification_docs_disk_file_id",
            "ufCrm8Certificationdocsdiskfileid",
        )
    )
    return ExistingCertificationDocsState(
        status=status,
        version=version,
        zip_url=zip_url,
        disk_file_id=disk_file_id,
    )


def load_lookup_catalog(path: Path) -> dict[str, Any]:
    payload = load_json_file(path)
    if isinstance(payload, dict):
        return {normalize_lookup_key(key): value for key, value in payload.items()}
    if not isinstance(payload, list):
        return {}
    result: dict[str, Any] = {}
    for row in payload:
        if not isinstance(row, dict):
            continue
        for key_name in ("onec_item_code", "sku", "item_ref"):
            key = clean_string(row.get(key_name))
            if key:
                result[normalize_lookup_key(key)] = row
    return result


def catalog_lookup(catalog: dict[str, Any], keys: list[str], value_key: str) -> str:
    for key in keys:
        if not key:
            continue
        value = catalog.get(normalize_lookup_key(key))
        if isinstance(value, dict):
            resolved = clean_string(value.get(value_key) or value.get(value_key.upper()))
            if resolved:
                return resolved
        elif isinstance(value, str):
            if value_key in {"barcode", "gtin", "certificate_id"}:
                return clean_string(value)
    return ""


def _certificate_from_mapping(value: Any, *, source: str = "catalog") -> CertificateCoverage:
    if isinstance(value, CertificateCoverage):
        return value
    if isinstance(value, str):
        certificate_id = clean_string(value)
        return CertificateCoverage(
            certificate_id=certificate_id,
            certificate_number=certificate_id,
            status="covered",
            eac_allowed=True,
            source=source,
        )
    if not isinstance(value, dict):
        return CertificateCoverage(source=source)
    status = clean_string(value.get("status") or value.get("certificate_status") or "covered")
    certificate_id = clean_string(
        value.get("certificate_id") or value.get("declaration_id") or value.get("number")
    )
    valid_to = clean_string(value.get("valid_to") or value.get("certificate_valid_to"))
    file_id = clean_string(value.get("file_id") or value.get("certificate_file_id"))
    covered_skus = split_sku_list(value.get("covered_skus") or value.get("sku"))
    return CertificateCoverage(
        item_id=clean_string(value.get("item_id") or value.get("certificate_item_id")),
        certificate_id=certificate_id,
        certificate_number=clean_string(value.get("number") or certificate_id),
        status=status,
        valid_to=valid_to,
        file_id=file_id,
        eac_allowed=_bool_value(value.get("eac", status in ACTIVE_CERTIFICATE_STATUSES)),
        covered_skus=covered_skus,
        source=source,
    )


def certificate_candidates_from_catalog(
    catalog: dict[str, Any], keys: list[str]
) -> list[CertificateCoverage]:
    candidates: list[CertificateCoverage] = []
    for key in keys:
        if not key:
            continue
        value = catalog.get(normalize_lookup_key(key))
        if not value:
            continue
        if isinstance(value, list):
            candidates.extend(_certificate_from_mapping(item) for item in value)
        else:
            candidates.append(_certificate_from_mapping(value))
    return candidates


def resolve_certificate(candidates: list[CertificateCoverage]) -> CertificateResolution:
    if not candidates:
        return CertificateResolution(blockers=(BLOCK_MISSING_CERTIFICATE,))

    active: list[CertificateCoverage] = []
    inactive_blockers: list[str] = []
    for candidate in candidates:
        status = clean_string(candidate.status).casefold()
        if status not in ACTIVE_CERTIFICATE_STATUSES:
            inactive_blockers.append(BLOCK_UNVERIFIED_CERTIFICATE)
            continue
        if _is_expired(candidate.valid_to):
            inactive_blockers.append(BLOCK_EXPIRED_CERTIFICATE)
            continue
        if not candidate.file_id:
            inactive_blockers.append(BLOCK_MISSING_CERTIFICATE_FILE)
            continue
        active.append(candidate)

    if len(active) > 1:
        return CertificateResolution(
            status="ambiguous",
            blockers=(BLOCK_AMBIGUOUS_CERTIFICATE,),
        )
    if len(active) == 1:
        return CertificateResolution(status="covered", certificate=active[0])
    blockers = tuple(dict.fromkeys(inactive_blockers or [BLOCK_MISSING_CERTIFICATE]))
    return CertificateResolution(status="blocked", blockers=blockers)


def _certificate_process_mapping(mapping: dict[str, Any]) -> dict[str, Any]:
    return _nested_mapping(mapping, "certificate_process")


def _product_passport_process_mapping(mapping: dict[str, Any]) -> dict[str, Any]:
    return _nested_mapping(mapping, "product_passport_process")


def _certificate_from_bitrix_item(
    item: dict[str, Any],
    *,
    field_map: dict[str, Any],
    enum_map: dict[str, Any],
) -> CertificateCoverage:
    title = clean_string(item.get("title") or item.get("TITLE"))
    status = _enum_status(
        _logical_field_value(item, field_map, "verification_status"),
        enum_map.get("verification_status") or {},
        labels={
            "Проверено": "verified",
            "Найдено": "found",
            "Требует обновления": "needs_update",
            "Истекло": "expired",
        },
        default="verified" if title else "missing",
    )
    valid_to = clean_string(
        _logical_field_value(item, field_map, "valid_to", "ufCrm7_1774687908027")
    )
    file_value = _logical_field_value(item, field_map, "file", "ufCrm7_1774688026350")
    certificate_number = clean_string(_logical_field_value(item, field_map, "declaration_number"))
    if not certificate_number:
        match = re.search(r"ЕАЭС\s+N\s+[^·,\n]+", title, re.IGNORECASE)
        certificate_number = match.group(0).strip() if match else title
    covered_skus = split_sku_list(_logical_field_value(item, field_map, "covered_skus"))
    return CertificateCoverage(
        item_id=clean_string(item.get("id") or item.get("ID")),
        certificate_id=certificate_number,
        certificate_number=certificate_number,
        status=status,
        valid_to=valid_to,
        file_id=_file_id(file_value),
        eac_allowed=_bool_value(_logical_field_value(item, field_map, "eac_allowed")),
        covered_skus=covered_skus,
        source="bitrix",
    )


def _product_passport_from_bitrix_item(
    item: dict[str, Any],
    *,
    field_map: dict[str, Any],
) -> ProductPassport:
    return ProductPassport(
        item_id=clean_string(item.get("id") or item.get("ID")),
        sku=clean_string(_logical_field_value(item, field_map, "sku", "ufCrm5_1774619613327")),
        onec_item_code=clean_string(
            _logical_field_value(item, field_map, "onec_item_code", "ufCrm5_1774619651702")
        ),
        trade_name=clean_string(
            _logical_field_value(item, field_map, "trade_name", "ufCrm5_1774619575350")
        ),
        tnved=clean_string(_logical_field_value(item, field_map, "tnved", "ufCrm5_1774619816679")),
        manufacturer=clean_string(_logical_field_value(item, field_map, "manufacturer")),
        product_series=clean_string(_logical_field_value(item, field_map, "product_series")),
    )


def load_bitrix_label_sources(
    *,
    client: ProcurementLabelsBitrixClient | None,
    mapping: dict[str, Any],
) -> BitrixLabelSources:
    if client is None:
        return BitrixLabelSources()

    passport_mapping = _product_passport_process_mapping(mapping)
    passport_entity_type_id = int(
        passport_mapping.get("entity_type_id") or PRODUCT_PASSPORT_ENTITY_TYPE_ID
    )
    passport_field_map = passport_mapping.get("field_map") or {}
    product_passports_by_sku: dict[str, ProductPassport] = {}
    try:
        for item in client.list_items(entity_type_id=passport_entity_type_id, limit=5000):
            passport = _product_passport_from_bitrix_item(item, field_map=passport_field_map)
            sku_key = normalize_lookup_key(passport.sku)
            if sku_key and sku_key not in product_passports_by_sku:
                product_passports_by_sku[sku_key] = passport
    except Exception:
        product_passports_by_sku = {}

    certificate_mapping = _certificate_process_mapping(mapping)
    certificate_entity_type_id = int(
        certificate_mapping.get("entity_type_id") or CERTIFICATE_ENTITY_TYPE_ID
    )
    certificate_field_map = certificate_mapping.get("field_map") or {}
    certificate_enum_map = certificate_mapping.get("enum_map") or {}
    certificates_by_sku: dict[str, list[CertificateCoverage]] = {}
    try:
        for item in client.list_items(entity_type_id=certificate_entity_type_id, limit=1000):
            certificate = _certificate_from_bitrix_item(
                item,
                field_map=certificate_field_map,
                enum_map=certificate_enum_map,
            )
            for sku in certificate.covered_skus:
                key = normalize_lookup_key(sku)
                if key:
                    certificates_by_sku.setdefault(key, []).append(certificate)
    except Exception:
        certificates_by_sku = {}

    return BitrixLabelSources(
        certificate_catalog=certificates_by_sku,
        product_passports=product_passports_by_sku,
    )


def fetch_onec_supplier_order_lines(engine: Engine, onec_number: str) -> list[dict[str, Any]]:
    candidates = source_number_candidates(onec_number)
    if not candidates:
        return []
    params = {f"number_{idx}": value for idx, value in enumerate(candidates)}
    placeholders = ", ".join(f":number_{idx}" for idx in range(len(candidates)))
    query = text(f"""
        WITH target_doc AS (
          SELECT TOP 1 doc._IDRRef
          FROM dbo._Document133 AS doc WITH (NOLOCK)
          WHERE LTRIM(RTRIM(doc._Number)) IN ({placeholders})
          ORDER BY doc._Date_Time DESC
        )
        SELECT
          LTRIM(RTRIM(doc._Number)) AS order_number,
          vt._LineNo2516 AS line_no,
          CONVERT(varchar(34), item._IDRRef, 1) AS item_ref_hex,
          item._Code AS onec_item_code,
          item._Description AS item_name,
          item._Fld836 AS article_1c,
          item._Fld9945 AS sku,
          barcode._Fld6984 AS barcode,
          unit._Description AS unit,
          vt._Fld2520 AS quantity,
          vt._Fld2529 AS price,
          vt._Fld2526 AS amount
        FROM target_doc
        JOIN dbo._Document133 AS doc WITH (NOLOCK)
          ON doc._IDRRef = target_doc._IDRRef
        JOIN dbo._Document133_VT2515 AS vt WITH (NOLOCK)
          ON vt._Document133_IDRRef = doc._IDRRef
        LEFT JOIN dbo._Reference62 AS item WITH (NOLOCK)
          ON item._IDRRef = vt._Fld2523RRef
        LEFT JOIN dbo._Reference41 AS unit WITH (NOLOCK)
          ON unit._IDRRef = vt._Fld2517RRef
        OUTER APPLY (
          SELECT TOP 1
            LTRIM(RTRIM(barcode_row._Fld6984)) AS _Fld6984
          FROM dbo._InfoRg6983 AS barcode_row WITH (NOLOCK)
          WHERE barcode_row._Fld6985_RRRef = item._IDRRef
            AND LTRIM(RTRIM(barcode_row._Fld6984)) <> ''
          ORDER BY barcode_row._Fld6984 ASC
        ) AS barcode
        ORDER BY vt._LineNo2516 ASC
        """)
    with engine.connect() as conn:
        return [dict(row) for row in conn.execute(query, params).mappings()]


def build_preview_from_sources(
    *,
    item_id: str,
    entity_type_id: int,
    bitrix_item: dict[str, Any],
    onec_lines: list[dict[str, Any]],
    mapping: dict[str, Any],
    barcode_catalog: dict[str, Any] | None = None,
    certificate_catalog: dict[str, Any] | None = None,
    product_passports: dict[str, ProductPassport] | None = None,
    existing_label_state: ExistingLabelState | None = None,
) -> ProcurementLabelOrderPreview:
    barcode_catalog = barcode_catalog or {}
    certificate_catalog = certificate_catalog or {}
    product_passports = product_passports or {}
    title = clean_string(bitrix_item.get("title") or bitrix_item.get("TITLE"))
    contour = item_procurement_contour(bitrix_item, mapping)
    onec_number = item_onec_number(bitrix_item, mapping)
    existing_label_state = existing_label_state or _existing_label_state(bitrix_item, mapping)
    global_blockers: list[str] = []
    if contour != "ved_import":
        global_blockers.append(BLOCK_NON_VED)
    if not onec_number:
        global_blockers.append("Не найден номер 1С заказа в карточке")
    if onec_number and not onec_lines:
        global_blockers.append(f"Не найдены строки заказа поставщику {onec_number} в 1С")

    rows: list[ProcurementLabelRow] = []
    for raw in onec_lines:
        onec_item_code = clean_string(raw.get("onec_item_code"))
        item_name = clean_string(raw.get("item_name"))
        article_1c = clean_string(raw.get("article_1c") or raw.get("article"))
        sku = clean_string(raw.get("sku"))
        item_ref_hex = clean_string(raw.get("item_ref_hex"))
        keys = [onec_item_code, sku, article_1c, item_ref_hex, item_name]
        catalog_barcode = catalog_lookup(barcode_catalog, keys, "barcode")
        catalog_gtin = catalog_lookup(barcode_catalog, keys, "gtin")
        onec_barcode = clean_string(raw.get("barcode") or raw.get("gtin"))
        barcode = catalog_barcode or catalog_gtin or onec_barcode
        barcode_source = (
            "catalog_barcode"
            if catalog_barcode
            else "catalog_gtin" if catalog_gtin else "1c_internal" if onec_barcode else ""
        )
        certificate_candidates = certificate_candidates_from_catalog(certificate_catalog, keys)
        certificate_resolution = resolve_certificate(certificate_candidates)
        certificate = certificate_resolution.certificate
        passport = product_passports.get(normalize_lookup_key(sku))
        if passport is None and article_1c:
            passport = product_passports.get(normalize_lookup_key(article_1c))
        label_warnings: list[str] = []
        if passport is None:
            label_warnings.append(WARNING_MISSING_PRODUCT_PASSPORT)
        else:
            if not passport.trade_name:
                label_warnings.append(WARNING_MISSING_TRADE_NAME)
            if not passport.tnved:
                label_warnings.append(WARNING_MISSING_TNVED)
        blockers: list[str] = []
        if not onec_item_code:
            blockers.append(BLOCK_MISSING_ONEC)
        if not sku:
            blockers.append(BLOCK_MISSING_SKU)
        if not barcode:
            blockers.append(BLOCK_MISSING_BARCODE)
        blockers.extend(certificate_resolution.blockers)
        rows.append(
            ProcurementLabelRow(
                line_no=int(raw.get("line_no") or 0),
                onec_item_code=onec_item_code,
                item_name=item_name,
                article_1c=article_1c,
                sku=sku,
                barcode=barcode,
                barcode_source=barcode_source,
                unit=clean_string(raw.get("unit")) or "шт",
                quantity=_decimal(raw.get("quantity")),
                price=_decimal(raw.get("price")) if raw.get("price") is not None else None,
                amount=_decimal(raw.get("amount")) if raw.get("amount") is not None else None,
                certificate_id=certificate.certificate_id if certificate else "",
                certificate_item_id=certificate.item_id if certificate else "",
                certificate_number=certificate.certificate_number if certificate else "",
                certificate_valid_to=certificate.valid_to if certificate else "",
                certificate_file_id=certificate.file_id if certificate else "",
                certificate_status=certificate_resolution.status,
                eac_allowed=bool(certificate.eac_allowed) if certificate else False,
                product_passport_item_id=passport.item_id if passport else "",
                trade_name=passport.trade_name if passport else "",
                tnved=passport.tnved if passport else "",
                manufacturer=passport.manufacturer if passport else "",
                product_series=passport.product_series if passport else "",
                label_warnings=label_warnings,
                status="blocked" if blockers else "ready",
                blockers=blockers,
            )
        )

    blockers = global_blockers + [
        f"строка {row.line_no}: {message}" for row in rows for message in row.blockers
    ]
    blocked = bool(blockers)
    return ProcurementLabelOrderPreview(
        item_id=item_id,
        entity_type_id=entity_type_id,
        onec_number=onec_number,
        title=title,
        contour=contour,
        status="blocked" if blocked else existing_label_state.status or "draft",
        ready=not blocked,
        blocked=blocked,
        blockers=blockers,
        rows=rows,
        artifact_version=existing_label_state.version,
        zip_url=existing_label_state.zip_url or None,
        disk_file_id=existing_label_state.disk_file_id or None,
    )


def _safe_slug(value: str) -> str:
    text_value = clean_string(value) or "order"
    return re.sub(r"[^0-9A-Za-zА-Яа-я._-]+", "-", text_value).strip("-") or "order"


def next_artifact_version(artifact_dir: Path, onec_number: str) -> int:
    slug = _safe_slug(onec_number)
    pattern = re.compile(rf"^ved-labels-{re.escape(slug)}-v(\d+)\.zip$")
    max_version = 0
    if artifact_dir.exists():
        for path in artifact_dir.glob(f"ved-labels-{slug}-v*.zip"):
            match = pattern.match(path.name)
            if match:
                max_version = max(max_version, int(match.group(1)))
    return max_version + 1


def next_certification_package_version(artifact_dir: Path, onec_number: str) -> int:
    slug = _safe_slug(onec_number)
    pattern = re.compile(rf"^ved-certification-docs-{re.escape(slug)}-v(\d+)\.zip$")
    max_version = 0
    if artifact_dir.exists():
        for path in artifact_dir.glob(f"ved-certification-docs-{slug}-v*.zip"):
            match = pattern.match(path.name)
            if match:
                max_version = max(max_version, int(match.group(1)))
    return max_version + 1


def _font(size: int, *, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for path in candidates:
        if path and Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def _trim_text(
    draw: ImageDraw.ImageDraw, text_value: str, font: ImageFont.ImageFont, width: int
) -> str:
    value = clean_string(text_value)
    if draw.textlength(value, font=font) <= width:
        return value
    suffix = "..."
    while value and draw.textlength(value + suffix, font=font) > width:
        value = value[:-1]
    return value + suffix if value else suffix


def _ean13_check_digit(first_12: str) -> str:
    total = 0
    for index, digit in enumerate(first_12):
        total += int(digit) * (1 if index % 2 == 0 else 3)
    return str((10 - (total % 10)) % 10)


def _ean13_bits(value: str) -> str | None:
    digits = re.sub(r"\D+", "", value)
    if len(digits) == 12:
        digits += _ean13_check_digit(digits)
    if len(digits) != 13:
        return None
    if _ean13_check_digit(digits[:12]) != digits[-1]:
        return None
    parity = EAN_PARITY[digits[0]]
    bits = "101"
    for digit, mode in zip(digits[1:7], parity, strict=True):
        bits += EAN_LEFT_ODD[digit] if mode == "O" else EAN_LEFT_EVEN[digit]
    bits += "01010"
    for digit in digits[7:]:
        bits += EAN_RIGHT[digit]
    bits += "101"
    return bits


def _draw_barcode(
    draw: ImageDraw.ImageDraw,
    *,
    barcode: str,
    left: int,
    top: int,
    width: int,
    height: int,
) -> None:
    bits = _ean13_bits(barcode)
    if not bits:
        for index in range(44):
            if (hash((barcode, index)) & 1) == 0:
                x0 = left + round(index * width / 44)
                x1 = left + round((index + 0.55) * width / 44)
                draw.rectangle((x0, top, x1, top + height), fill="black")
        return
    module = max(1, width // len(bits))
    barcode_width = module * len(bits)
    offset = left + max(0, (width - barcode_width) // 2)
    for index, bit in enumerate(bits):
        if bit == "1":
            x0 = offset + index * module
            draw.rectangle((x0, top, x0 + module - 1, top + height), fill="black")


def render_label_png(row: ProcurementLabelRow, path: Path) -> None:
    image = Image.new("RGB", (LABEL_WIDTH_PX, LABEL_HEIGHT_PX), "white")
    draw = ImageDraw.Draw(image)
    margin = 28
    title_font = _font(30, bold=True)
    text_font = _font(24)
    small_font = _font(20)
    eac_font = _font(44, bold=True)
    max_width = LABEL_WIDTH_PX - margin * 2

    draw.rectangle((0, 0, LABEL_WIDTH_PX - 1, LABEL_HEIGHT_PX - 1), outline="black", width=3)
    draw.text(
        (margin, 22),
        _trim_text(draw, row.sku, title_font, max_width),
        fill="black",
        font=title_font,
    )
    product_title = row.trade_name or row.item_name
    draw.text(
        (margin, 62),
        _trim_text(draw, product_title, text_font, max_width),
        fill="black",
        font=text_font,
    )
    details = f"1C: {row.onec_item_code}"
    if row.tnved:
        details += f" | TNVED: {row.tnved}"
    draw.text(
        (margin, 98),
        _trim_text(draw, details, small_font, max_width),
        fill="black",
        font=small_font,
    )
    if row.product_series or row.manufacturer:
        series = " | ".join(value for value in (row.product_series, row.manufacturer) if value)
        draw.text(
            (margin, 124),
            _trim_text(draw, series, small_font, max_width - 130),
            fill="black",
            font=small_font,
        )
    if row.eac_allowed:
        draw.text((LABEL_WIDTH_PX - 120, 102), "EAC", fill="black", font=eac_font)

    barcode_top = 168
    _draw_barcode(
        draw,
        barcode=row.barcode,
        left=margin,
        top=barcode_top,
        width=max_width,
        height=140,
    )
    barcode_text = _trim_text(draw, row.barcode, title_font, max_width)
    text_width = draw.textlength(barcode_text, font=title_font)
    draw.text(
        ((LABEL_WIDTH_PX - text_width) / 2, barcode_top + 148),
        barcode_text,
        fill="black",
        font=title_font,
    )
    draw.text(
        (margin, LABEL_HEIGHT_PX - 40),
        _trim_text(draw, row.certificate_id, small_font, max_width),
        fill="black",
        font=small_font,
    )
    image.save(path, format="PNG", dpi=(LABEL_DPI, LABEL_DPI))


def write_register_xlsx(preview: ProcurementLabelOrderPreview, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "label-register"
    headers = [
        "line_no",
        "onec_order",
        "onec_item_code",
        "article_1c",
        "sku",
        "barcode",
        "barcode_source",
        "item_name",
        "trade_name",
        "tnved",
        "qty",
        "unit",
        "certificate_id",
        "certificate_item_id",
        "certificate_number",
        "certificate_valid_to",
        "certificate_file_id",
        "certificate_status",
        "eac_on_label",
        "label_warnings",
        "status",
        "blockers",
    ]
    ws.append(headers)
    for row in preview.rows:
        ws.append(
            [
                row.line_no,
                preview.onec_number,
                row.onec_item_code,
                row.article_1c,
                row.sku,
                row.barcode,
                row.barcode_source,
                row.item_name,
                row.trade_name,
                row.tnved,
                float(row.quantity),
                row.unit,
                row.certificate_id,
                row.certificate_item_id,
                row.certificate_number,
                row.certificate_valid_to,
                row.certificate_file_id,
                row.certificate_status,
                "yes" if row.eac_allowed else "no",
                "; ".join(row.label_warnings),
                row.status,
                "; ".join(row.blockers),
            ]
        )
    for column_cells in ws.columns:
        length = max(len(clean_string(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 48)
    wb.save(path)


def write_factory_readme(preview: ProcurementLabelOrderPreview, path: Path) -> None:
    lines = [
        f"Партия: {preview.onec_number}",
        f"Карточка Bitrix: {preview.item_id}",
        "Формат этикетки: 58x40 мм, PNG, 300 dpi.",
        "На этикетке нет Честного знака и РСТ.",
        "EAC выводится только для строк, где товар связан с подтвержденной ДС/сертификатом.",
        "Печатать по одному макету на SKU; количество брать из label-register.xlsx.",
        "",
        "Файлы:",
        "- labels-preview/*.png — макеты товарных этикеток",
        "- label-register.xlsx — соответствие 1С -> артикул -> SKU -> штрихкод -> ДС -> строка заказа",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_artifacts(
    preview: ProcurementLabelOrderPreview,
    *,
    artifact_dir: Path,
    base_url: str = "",
    version: int | None = None,
) -> tuple[int, Path, str]:
    if preview.blocked:
        raise ValueError("Cannot generate labels while preview is blocked")
    artifact_dir.mkdir(parents=True, exist_ok=True)
    resolved_version = version or next_artifact_version(artifact_dir, preview.onec_number)
    slug = _safe_slug(preview.onec_number)
    work_dir = artifact_dir / f"{slug}-v{resolved_version}"
    labels_dir = work_dir / "labels-preview"
    labels_dir.mkdir(parents=True, exist_ok=True)

    for row in preview.rows:
        filename = f"line-{row.line_no:03d}-{_safe_slug(row.sku or row.onec_item_code)}.png"
        render_label_png(row, labels_dir / filename)
    write_register_xlsx(preview, work_dir / "label-register.xlsx")
    write_factory_readme(preview, work_dir / "factory-labels-readme.txt")

    zip_path = artifact_dir / f"ved-labels-{slug}-v{resolved_version}.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(work_dir.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(work_dir).as_posix())

    if base_url:
        zip_url = f"{base_url.rstrip('/')}/api/procurement-labels/artifacts/{quote(zip_path.name)}"
    else:
        zip_url = zip_path.as_posix()
    return resolved_version, zip_path, zip_url


def _certification_row_missing_fields(row: dict[str, Any]) -> list[str]:
    return [field for field in CERTIFICATION_REQUIRED_FIELDS if not clean_string(row.get(field))]


def _certification_review_status(row: dict[str, Any], missing_fields: list[str]) -> str:
    if "sku_1c" in missing_fields:
        return "missing_sku_1c"
    if "gtin_ean13" in missing_fields and (
        "gost_r_ds_number" in missing_fields or "gost_r_ds_covers_sku" in missing_fields
    ):
        return "needs_gtin_and_gost_r_ds"
    if "gost_r_ds_number" in missing_fields or "gost_r_ds_covers_sku" in missing_fields:
        return "needs_gost_r_ds"
    if "gtin_ean13" in missing_fields:
        return "needs_gtin"
    if missing_fields:
        return "needs_product_properties"
    return "ready_for_certifier"


def _certification_next_action(row: dict[str, Any], missing_fields: list[str]) -> str:
    if "sku_1c" in missing_fields:
        return "Заполнить отдельное поле SKU в 1С; числовой Артикул не использовать как SKU."
    if "gtin_ean13" in missing_fields and (
        "gost_r_ds_number" in missing_fields or "gost_r_ds_covers_sku" in missing_fields
    ):
        return "Заказать GTIN/EAN-13 и передать строку сертификатору для покрытия новой ДС ГОСТ Р."
    if "gtin_ean13" in missing_fields:
        return "Заказать GTIN/EAN-13 и вернуть номер в мастер-реестр."
    if "gost_r_ds_number" in missing_fields or "gost_r_ds_covers_sku" in missing_fields:
        return "Получить номер новой ДС ГОСТ Р и подтвердить покрытие SKU."
    if missing_fields:
        return "Дозаполнить свойства товара перед финальной подачей."
    return "Можно передавать сертификатору на финальную проверку."


def _certification_master_rows(preview: ProcurementLabelOrderPreview) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in preview.rows:
        master_row: dict[str, Any] = {
            "line_no": row.line_no,
            "onec_order": preview.onec_number,
            "onec_item_code": row.onec_item_code,
            "article_1c": row.article_1c,
            "sku_1c": row.sku,
            "item_name_1c": row.item_name,
            "qty": float(row.quantity),
            "unit": row.unit,
            "barcode_1c": row.barcode,
            "barcode_source": row.barcode_source,
            "generated_sku": "",
            "generation_status": "",
            "external_sku_candidate": "",
            "trade_name_family": row.trade_name,
            "compatibility": row.product_series,
            "capacity_mah": "",
            "voltage_v": "",
            "energy_wh": "",
            "dim": "",
            "bms": "",
            "connector": "",
            "tnved": row.tnved or "8507600000",
            "un_code": "UN3480",
            "current_eaeu_ds_number": row.certificate_number,
            "new_gost_r_ds_number": "",
            "gost_r_ds_covers_sku": "",
            "gtin_ean13": "",
        }
        missing_fields = _certification_row_missing_fields(master_row)
        master_row["missing_fields"] = "; ".join(missing_fields)
        master_row["review_status"] = _certification_review_status(master_row, missing_fields)
        master_row["next_action"] = _certification_next_action(master_row, missing_fields)
        rows.append(master_row)
    return rows


def _certification_column_title(column: str) -> str:
    return CERTIFICATION_COLUMN_TITLES.get(column, column)


def _certification_missing_fields_title(value: Any) -> str:
    fields = [clean_string(item) for item in clean_string(value).split(";") if clean_string(item)]
    return "; ".join(_certification_column_title(field) for field in fields)


def _certification_display_cell(column: str, value: Any) -> Any:
    if column == "missing_fields":
        return _certification_missing_fields_title(value)
    if column == "review_status":
        return CERTIFICATION_REVIEW_STATUS_TITLES.get(clean_string(value), value)
    if column == "barcode_source":
        return CERTIFICATION_BARCODE_SOURCE_TITLES.get(clean_string(value), value)
    return value


def _write_simple_table_xlsx(
    path: Path, *, title: str, columns: list[str], rows: list[dict[str, Any]]
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append([_certification_column_title(column) for column in columns])
    for row in rows:
        ws.append([_certification_display_cell(column, row.get(column, "")) for column in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        length = max(len(clean_string(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 60)
    wb.save(path)


def _write_certification_master_xlsx(rows: list[dict[str, Any]], path: Path) -> None:
    columns = [
        "line_no",
        "onec_order",
        "onec_item_code",
        "article_1c",
        "sku_1c",
        "item_name_1c",
        "qty",
        "unit",
        "barcode_1c",
        "barcode_source",
        "generated_sku",
        "generation_status",
        "external_sku_candidate",
        "trade_name_family",
        "compatibility",
        "capacity_mah",
        "voltage_v",
        "energy_wh",
        "dim",
        "bms",
        "connector",
        "tnved",
        "un_code",
        "current_eaeu_ds_number",
        "new_gost_r_ds_number",
        "gost_r_ds_covers_sku",
        "gtin_ean13",
        "missing_fields",
        "review_status",
        "next_action",
    ]
    _write_simple_table_xlsx(
        path,
        title=CERTIFICATION_SHEET_TITLES["master"],
        columns=columns,
        rows=rows,
    )


def _write_gtin_order_xlsx(rows: list[dict[str, Any]], path: Path) -> None:
    gtin_rows = [
        {
            "line_no": row["line_no"],
            "onec_item_code": row["onec_item_code"],
            "article_1c": row["article_1c"],
            "sku_1c": row["sku_1c"],
            "item_name_1c": row["item_name_1c"],
            "qty": row["qty"],
            "barcode_1c": row["barcode_1c"],
            "gtin_ean13": row["gtin_ean13"],
            "gtin_action": "заказать GTIN/EAN-13" if not row["gtin_ean13"] else "проверить",
        }
        for row in rows
    ]
    _write_simple_table_xlsx(
        path,
        title=CERTIFICATION_SHEET_TITLES["gtin"],
        columns=[
            "line_no",
            "onec_item_code",
            "article_1c",
            "sku_1c",
            "item_name_1c",
            "qty",
            "barcode_1c",
            "gtin_ean13",
            "gtin_action",
        ],
        rows=gtin_rows,
    )


def _write_certification_checklist(path: Path) -> None:
    lines = [
        "# Чек-лист пакета для ДС ГОСТ Р",
        "",
        "Собрать и проверить перед передачей сертификатору:",
        "",
    ]
    lines.extend(f"- {item}" for item in CERTIFICATION_DOCUMENT_CHECKLIST)
    lines.extend(
        [
            "",
            "Важно: TradeName/family_formula используется для декларации, инвойса и паспорта товара.",
            "SKU остается отдельным идентификатором товара; числовой Артикул 1С SKU не заменяет.",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_broker_request(
    preview: ProcurementLabelOrderPreview, rows: list[dict[str, Any]], path: Path
) -> None:
    sku_count = sum(1 for row in rows if clean_string(row.get("sku_1c")))
    missing_count = sum(1 for row in rows if clean_string(row.get("missing_fields")))
    lines = [
        f"Партия: {preview.onec_number}",
        f"Карточка Bitrix: {preview.item_id}",
        f"Строк: {len(rows)}",
        f"SKU с заполненным полем SKU 1С: {sku_count}",
        f"Строк с недостающими данными: {missing_count}",
        "",
        "Просьба к брокеру/сертификатору:",
        "- подтвердить, можно ли закрыть партию одной семейной ДС ГОСТ Р по TradeName/family_formula;",
        "- подтвердить перечень SKU, которые должны войти в область покрытия ДС;",
        "- подтвердить, какие GTIN/EAN-13 обязательны до регистрации/маркировки;",
        "- вернуть точную формулировку продукции для декларации и инвойса;",
        "- указать, какие образцы и протоколы испытаний нужны по выбранной схеме.",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _find_existing_master_register(
    *, preview: ProcurementLabelOrderPreview, artifact_dir: Path
) -> Path | None:
    slug = _safe_slug(preview.onec_number)
    for work_dir in sorted(artifact_dir.glob(f"{slug}-v*"), reverse=True):
        if not work_dir.is_dir():
            continue
        candidates = sorted(
            [
                *work_dir.glob("*мастер*таблица*.xlsx"),
                *work_dir.glob("*master-register.xlsx"),
            ]
        )
        if candidates:
            return candidates[-1]
    return None


def certification_package_blockers(preview: ProcurementLabelOrderPreview) -> list[str]:
    blockers: list[str] = []
    if preview.contour != "ved_import":
        blockers.append(BLOCK_NON_VED)
    if not preview.onec_number:
        blockers.append("Не найден номер 1С заказа в карточке")
    if not preview.rows:
        blockers.append(BLOCK_MISSING_CERTIFICATION_ROWS)
    return blockers


def build_certification_docs_artifacts(
    preview: ProcurementLabelOrderPreview,
    *,
    artifact_dir: Path,
    base_url: str = "",
    version: int | None = None,
) -> tuple[int, Path, str, int, int]:
    blockers = certification_package_blockers(preview)
    if blockers:
        raise ValueError("Cannot generate certification docs package: " + "; ".join(blockers))
    artifact_dir.mkdir(parents=True, exist_ok=True)
    resolved_version = version or next_certification_package_version(
        artifact_dir, preview.onec_number
    )
    slug = _safe_slug(preview.onec_number)
    work_dir = artifact_dir / f"{slug}-certification-docs-v{resolved_version}"
    work_dir.mkdir(parents=True, exist_ok=True)

    master_rows = _certification_master_rows(preview)
    existing_master = _find_existing_master_register(preview=preview, artifact_dir=artifact_dir)
    if existing_master is not None:
        shutil.copy2(existing_master, work_dir / CERTIFICATION_MASTER_FILENAME)
    else:
        _write_certification_master_xlsx(master_rows, work_dir / CERTIFICATION_MASTER_FILENAME)
    _write_gtin_order_xlsx(master_rows, work_dir / "gtin-order-list.xlsx")
    _write_certification_checklist(work_dir / "certification-documents-checklist.md")
    _write_broker_request(preview, master_rows, work_dir / "broker-request.md")
    (work_dir / "package-readme.txt").write_text(
        "\n".join(
            [
                f"Партия: {preview.onec_number}",
                "Назначение: пакет для заказа GTIN/EAN-13 и подготовки ДС ГОСТ Р.",
                f"Файл {CERTIFICATION_MASTER_FILENAME} является рабочим реестром для декларации.",
                "Файл gtin-order-list.xlsx содержит SKU, по которым нужно заказать GTIN.",
                "Файл broker-request.md можно отправить брокеру/сертификатору как краткое поручение.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    zip_path = artifact_dir / f"ved-certification-docs-{slug}-v{resolved_version}.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(work_dir.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(work_dir).as_posix())

    if base_url:
        zip_url = f"{base_url.rstrip('/')}/api/procurement-labels/artifacts/{quote(zip_path.name)}"
    else:
        zip_url = zip_path.as_posix()
    gtin_rows = sum(1 for row in master_rows if not clean_string(row.get("gtin_ean13")))
    missing_rows = sum(1 for row in master_rows if clean_string(row.get("missing_fields")))
    return resolved_version, zip_path, zip_url, gtin_rows, missing_rows


def upload_zip_to_disk(
    *,
    settings: Settings,
    zip_path: Path,
) -> BitrixFileResult:
    webhook_url = (
        settings.procurement_labels_bitrix_webhook_url
        or settings.procurement_bitrix_webhook_url
        or settings.bitrix_box_webhook_base
    )
    if not webhook_url:
        raise RuntimeError("Bitrix webhook for procurement labels is not configured")
    if not settings.procurement_labels_bitrix_root_folder_id:
        raise RuntimeError("Bitrix Disk folder for procurement labels is not configured")
    client = BitrixDiskClient(
        webhook_url,
        timeout=int(settings.procurement_labels_bitrix_rest_timeout_seconds),
    )
    try:
        result = client.upload_file(
            int(settings.procurement_labels_bitrix_root_folder_id),
            filename=zip_path.name,
            content=zip_path.read_bytes(),
        )
    except BitrixDiskError as exc:
        raise RuntimeError(str(exc)) from exc
    file_id = clean_string(result.get("ID") or result.get("id") or result.get("fileId"))
    url = clean_string(
        result.get("DETAIL_URL")
        or result.get("detailUrl")
        or result.get("DOWNLOAD_URL")
        or result.get("downloadUrl")
    )
    if file_id and not url:
        try:
            file_payload = client.call("disk.file.get", {"id": file_id})
            file_result = file_payload.get("result") or {}
            if isinstance(file_result, dict):
                url = clean_string(
                    file_result.get("DETAIL_URL")
                    or file_result.get("detailUrl")
                    or file_result.get("DOWNLOAD_URL")
                    or file_result.get("downloadUrl")
                )
        except BitrixDiskError:
            url = ""
    if not file_id or not url:
        raise RuntimeError("Bitrix Disk did not return file id and link for labels ZIP")
    return BitrixFileResult(file_id=file_id, url=url)


def label_status_fields(
    mapping: dict[str, Any],
    *,
    status: str,
    version: int | None = None,
    zip_url: str | None = None,
    disk_file_id: str | None = None,
    errors: list[str] | None = None,
) -> dict[str, Any]:
    fields: dict[str, Any] = {}
    field_map = mapping.get("field_map") or {}
    enum_map = mapping.get("enum_map") or {}
    status_field = clean_string(field_map.get("label_generation_status"))
    if status_field:
        enum_id = clean_string((enum_map.get("label_generation_status") or {}).get(status))
        fields[crm_item_rest_field_name(status_field)] = enum_id or status
    version_field = clean_string(field_map.get("label_generation_version"))
    if version_field and version is not None:
        fields[crm_item_rest_field_name(version_field)] = version
    zip_field = clean_string(field_map.get("label_generation_zip_url"))
    if zip_field and zip_url:
        fields[crm_item_rest_field_name(zip_field)] = zip_url
    disk_file_field = clean_string(field_map.get("label_generation_disk_file_id"))
    if disk_file_field and disk_file_id:
        fields[crm_item_rest_field_name(disk_file_field)] = disk_file_id
    errors_field = clean_string(field_map.get("label_generation_errors"))
    if errors_field:
        fields[crm_item_rest_field_name(errors_field)] = "\n".join(errors or [])
    approved_at_field = clean_string(field_map.get("label_generation_approved_at"))
    if approved_at_field and status == "approved":
        fields[crm_item_rest_field_name(approved_at_field)] = datetime.now(UTC).isoformat()
    sent_at_field = clean_string(field_map.get("label_generation_sent_to_factory_at"))
    if sent_at_field and status == "sent_to_factory":
        fields[crm_item_rest_field_name(sent_at_field)] = datetime.now(UTC).isoformat()
    return fields


def certification_docs_status_fields(
    mapping: dict[str, Any],
    *,
    status: str,
    version: int | None = None,
    zip_url: str | None = None,
    disk_file_id: str | None = None,
    errors: list[str] | None = None,
) -> dict[str, Any]:
    fields: dict[str, Any] = {}
    field_map = mapping.get("field_map") or {}
    enum_map = mapping.get("enum_map") or {}
    status_field = clean_string(field_map.get("certification_docs_status"))
    if status_field:
        enum_id = clean_string((enum_map.get("certification_docs_status") or {}).get(status))
        fields[crm_item_rest_field_name(status_field)] = enum_id or status
    version_field = clean_string(field_map.get("certification_docs_version"))
    if version_field and version is not None:
        fields[crm_item_rest_field_name(version_field)] = version
    zip_field = clean_string(field_map.get("certification_docs_zip_url"))
    if zip_field and zip_url:
        fields[crm_item_rest_field_name(zip_field)] = zip_url
    disk_file_field = clean_string(field_map.get("certification_docs_disk_file_id"))
    if disk_file_field and disk_file_id:
        fields[crm_item_rest_field_name(disk_file_field)] = disk_file_id
    errors_field = clean_string(field_map.get("certification_docs_errors"))
    if errors_field:
        fields[crm_item_rest_field_name(errors_field)] = "\n".join(errors or [])
    generated_at_field = clean_string(field_map.get("certification_docs_generated_at"))
    if generated_at_field and status not in {"blocked"}:
        fields[crm_item_rest_field_name(generated_at_field)] = datetime.now(UTC).isoformat()
    return fields


def update_bitrix_label_status(
    *,
    client: ProcurementLabelsBitrixClient | None,
    entity_type_id: int,
    item_id: str,
    mapping: dict[str, Any],
    status: str,
    version: int | None = None,
    zip_url: str | None = None,
    disk_file_id: str | None = None,
    errors: list[str] | None = None,
) -> None:
    if client is None:
        return
    fields = label_status_fields(
        mapping,
        status=status,
        version=version,
        zip_url=zip_url,
        disk_file_id=disk_file_id,
        errors=errors,
    )
    client.update_item(entity_type_id=entity_type_id, item_id=item_id, fields=fields)


def update_bitrix_certification_docs_status(
    *,
    client: ProcurementLabelsBitrixClient | None,
    entity_type_id: int,
    item_id: str,
    mapping: dict[str, Any],
    status: str,
    version: int | None = None,
    zip_url: str | None = None,
    disk_file_id: str | None = None,
    errors: list[str] | None = None,
) -> None:
    if client is None:
        return
    fields = certification_docs_status_fields(
        mapping,
        status=status,
        version=version,
        zip_url=zip_url,
        disk_file_id=disk_file_id,
        errors=errors,
    )
    client.update_item(entity_type_id=entity_type_id, item_id=item_id, fields=fields)


def bitrix_client_from_settings(settings: Settings) -> ProcurementLabelsBitrixClient | None:
    webhook_url = (
        settings.procurement_labels_bitrix_webhook_url
        or settings.procurement_bitrix_webhook_url
        or settings.bitrix_box_webhook_base
    )
    if not webhook_url:
        return None
    return ProcurementLabelsBitrixClient(
        webhook_url,
        timeout=settings.procurement_labels_bitrix_rest_timeout_seconds,
    )


def onec_engine_from_settings(settings: Settings) -> Engine:
    if not settings.onec_database_url:
        raise HTTPException(status_code=500, detail="1C database URL is not configured")
    return build_onec_engine_from_settings()


def build_preview(
    item_id: str,
    *,
    settings: Settings | None = None,
    bitrix_client: ProcurementLabelsBitrixClient | None = None,
    onec_engine: Engine | None = None,
) -> ProcurementLabelOrderPreview:
    settings = settings or get_settings()
    mapping = load_mapping(settings)
    entity_type_id = int(
        ((mapping.get("process") or {}).get("entity_type_id"))
        or settings.procurement_labels_entity_type_id
    )
    client = bitrix_client or bitrix_client_from_settings(settings)
    if client is None:
        raise HTTPException(status_code=500, detail="Bitrix procurement webhook is not configured")
    item = client.get_item(entity_type_id=entity_type_id, item_id=item_id)
    onec_number = item_onec_number(item, mapping)
    engine = onec_engine or onec_engine_from_settings(settings)
    onec_lines = fetch_onec_supplier_order_lines(engine, onec_number) if onec_number else []
    barcode_catalog = load_lookup_catalog(Path(settings.procurement_labels_barcode_catalog_path))
    local_certificate_catalog = load_lookup_catalog(
        Path(settings.procurement_labels_certificate_catalog_path)
    )
    bitrix_sources = load_bitrix_label_sources(client=client, mapping=mapping)
    certificate_catalog = bitrix_sources.certificate_catalog or local_certificate_catalog
    return build_preview_from_sources(
        item_id=item_id,
        entity_type_id=entity_type_id,
        bitrix_item=item,
        onec_lines=onec_lines,
        mapping=mapping,
        barcode_catalog=barcode_catalog,
        certificate_catalog=certificate_catalog,
        product_passports=bitrix_sources.product_passports,
        existing_label_state=_existing_label_state(item, mapping),
    )


def generate_zip(
    item_id: str,
    *,
    settings: Settings | None = None,
    bitrix_client: ProcurementLabelsBitrixClient | None = None,
    onec_engine: Engine | None = None,
    base_url: str = "",
    dry_run: bool = False,
) -> ProcurementLabelGenerateResponse:
    settings = settings or get_settings()
    mapping = load_mapping(settings)
    client = bitrix_client or bitrix_client_from_settings(settings)
    preview = build_preview(
        item_id,
        settings=settings,
        bitrix_client=client,
        onec_engine=onec_engine,
    )
    if preview.blocked:
        if not dry_run:
            update_bitrix_label_status(
                client=client,
                entity_type_id=preview.entity_type_id,
                item_id=item_id,
                mapping=mapping,
                status="blocked",
                errors=preview.blockers,
            )
        return ProcurementLabelGenerateResponse(preview=preview, generated=False)

    version, zip_path, zip_url = build_artifacts(
        preview,
        artifact_dir=Path(settings.procurement_labels_artifact_dir),
        base_url=base_url,
    )
    disk_file = None if dry_run else upload_zip_to_disk(settings=settings, zip_path=zip_path)
    final_url = zip_url if dry_run else disk_file.url
    preview.artifact_version = version
    preview.zip_url = final_url
    preview.disk_file_id = disk_file.file_id if disk_file else None
    if not dry_run:
        update_bitrix_label_status(
            client=client,
            entity_type_id=preview.entity_type_id,
            item_id=item_id,
            mapping=mapping,
            status="draft",
            version=version,
            zip_url=final_url,
            disk_file_id=disk_file.file_id if disk_file else None,
            errors=[],
        )
    return ProcurementLabelGenerateResponse(
        preview=preview,
        generated=True,
        artifact_version=version,
        zip_filename=zip_path.name,
        zip_url=final_url,
        disk_file_id=disk_file.file_id if disk_file else None,
    )


def generate_certification_docs_zip(
    item_id: str,
    *,
    settings: Settings | None = None,
    bitrix_client: ProcurementLabelsBitrixClient | None = None,
    onec_engine: Engine | None = None,
    base_url: str = "",
    dry_run: bool = False,
) -> ProcurementCertificationDocsGenerateResponse:
    settings = settings or get_settings()
    mapping = load_mapping(settings)
    client = bitrix_client or bitrix_client_from_settings(settings)
    preview = build_preview(
        item_id,
        settings=settings,
        bitrix_client=client,
        onec_engine=onec_engine,
    )
    blockers = certification_package_blockers(preview)
    if blockers:
        if not dry_run:
            update_bitrix_certification_docs_status(
                client=client,
                entity_type_id=preview.entity_type_id,
                item_id=item_id,
                mapping=mapping,
                status="blocked",
                errors=blockers,
            )
        return ProcurementCertificationDocsGenerateResponse(
            preview=preview,
            generated=False,
            document_checklist=CERTIFICATION_DOCUMENT_CHECKLIST,
        )

    version, zip_path, zip_url, gtin_rows, missing_rows = build_certification_docs_artifacts(
        preview,
        artifact_dir=Path(settings.procurement_labels_artifact_dir),
        base_url=base_url,
    )
    disk_file = None if dry_run else upload_zip_to_disk(settings=settings, zip_path=zip_path)
    final_url = zip_url if dry_run else disk_file.url
    status = "needs_data" if missing_rows else "ready"
    if not dry_run:
        update_bitrix_certification_docs_status(
            client=client,
            entity_type_id=preview.entity_type_id,
            item_id=item_id,
            mapping=mapping,
            status=status,
            version=version,
            zip_url=final_url,
            disk_file_id=disk_file.file_id if disk_file else None,
            errors=[],
        )
    return ProcurementCertificationDocsGenerateResponse(
        preview=preview,
        generated=True,
        artifact_version=version,
        zip_filename=zip_path.name,
        zip_url=final_url,
        disk_file_id=disk_file.file_id if disk_file else None,
        gtin_rows=gtin_rows,
        missing_rows=missing_rows,
        document_checklist=CERTIFICATION_DOCUMENT_CHECKLIST,
    )


def approve_zip(
    item_id: str,
    *,
    settings: Settings | None = None,
    bitrix_client: ProcurementLabelsBitrixClient | None = None,
) -> tuple[int | None, str | None]:
    settings = settings or get_settings()
    mapping = load_mapping(settings)
    client = bitrix_client or bitrix_client_from_settings(settings)
    entity_type_id = int(
        ((mapping.get("process") or {}).get("entity_type_id"))
        or settings.procurement_labels_entity_type_id
    )
    if client is None:
        raise HTTPException(status_code=500, detail="Bitrix procurement webhook is not configured")
    item = client.get_item(entity_type_id=entity_type_id, item_id=item_id)
    state = _existing_label_state(item, mapping)
    if not state.version or not state.zip_url or not state.disk_file_id:
        raise HTTPException(
            status_code=409,
            detail="Сначала сформируйте ZIP и дождитесь загрузки файла в Bitrix Disk",
        )
    preview = build_preview(item_id, settings=settings, bitrix_client=client)
    if preview.blocked:
        update_bitrix_label_status(
            client=client,
            entity_type_id=entity_type_id,
            item_id=item_id,
            mapping=mapping,
            status="blocked",
            errors=preview.blockers,
        )
        raise HTTPException(
            status_code=409,
            detail="Нельзя утвердить макет: есть стоп-ошибки",
        )
    update_bitrix_label_status(
        client=client,
        entity_type_id=entity_type_id,
        item_id=item_id,
        mapping=mapping,
        status="approved",
        version=state.version,
        zip_url=state.zip_url,
        disk_file_id=state.disk_file_id,
        errors=[],
    )
    return state.version, state.zip_url


def send_zip_to_factory(
    item_id: str,
    *,
    settings: Settings | None = None,
    bitrix_client: ProcurementLabelsBitrixClient | None = None,
) -> tuple[int | None, str | None]:
    settings = settings or get_settings()
    mapping = load_mapping(settings)
    client = bitrix_client or bitrix_client_from_settings(settings)
    entity_type_id = int(
        ((mapping.get("process") or {}).get("entity_type_id"))
        or settings.procurement_labels_entity_type_id
    )
    if client is None:
        raise HTTPException(status_code=500, detail="Bitrix procurement webhook is not configured")
    item = client.get_item(entity_type_id=entity_type_id, item_id=item_id)
    state = _existing_label_state(item, mapping)
    if state.status != "approved":
        raise HTTPException(
            status_code=409,
            detail="Сначала утвердите версию ZIP, затем отмечайте отправку фабрике",
        )
    if not state.version or not state.zip_url or not state.disk_file_id:
        raise HTTPException(
            status_code=409,
            detail="В карточке нет утвержденной версии ZIP с файлом Bitrix Disk",
        )
    update_bitrix_label_status(
        client=client,
        entity_type_id=entity_type_id,
        item_id=item_id,
        mapping=mapping,
        status="sent_to_factory",
        version=state.version,
        zip_url=state.zip_url,
        disk_file_id=state.disk_file_id,
        errors=[],
    )
    return state.version, state.zip_url
